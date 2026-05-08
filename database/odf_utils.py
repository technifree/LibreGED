from odf.opendocument import load
from odf import text, table, draw, presentation
from odf import teletype
from zipfile import ZipFile
from io import BytesIO
import os
import sys
import config  # Importer config pour obtenir FILES_DIR


def odt_to_html(path):
    """Convertir un fichier ODT en HTML"""
    path = config.FILES_DIR / path  # Utiliser config.FILES_DIR pour obtenir le chemin correct
    doc = load(str(path))
    
    # Structure HTML de base
    html = "<html><head><meta charset='UTF-8'><title>Document ODT</title></head><body>"

    # Traitement des titres (H)
    for elem in doc.getElementsByType(text.H):
        level = elem.getAttribute("outlinelevel")
        content = teletype.extractText(elem)
        html += f"<h{level}>{content}</h{level}>\n"
    
    # Traitement des paragraphes (P)
    for elem in doc.getElementsByType(text.P):
        content = teletype.extractText(elem)
        html += f"<p>{content}</p>\n"

    html += "</body></html>"  # Fermeture des balises HTML

    return html

def ods_to_html(path):
    """Convertir un fichier ODS en HTML"""
    path = config.FILES_DIR / path  # Utiliser config.FILES_DIR
    doc = load(str(path))
    html = "<table border='1'>\n"
    for table_elem in doc.spreadsheet.getElementsByType(table.Table):
        for row in table_elem.getElementsByType(table.TableRow):
            html += "<tr>"
            for cell in row.getElementsByType(table.TableCell):
                cell_text = ""
                for p in cell.getElementsByType(text.P):
                    cell_text += teletype.extractText(p)
                html += f"<td>{cell_text}</td>"
            html += "</tr>\n"
    html += "</table>"
    return html

def odp_to_html(path):
    """Convertir un fichier ODP en HTML"""
    path = config.FILES_DIR / path  # Utiliser config.FILES_DIR
    doc = load(str(path))
    html = ""
    for slide in doc.presentation.getElementsByType(draw.Page):
        html += f"<h2>Slide: {slide.getAttribute('draw:name')}</h2>\n"
        for frame in slide.getElementsByType(draw.Frame):
            for box in frame.getElementsByType(draw.TextBox):
                for p in box.getElementsByType(text.P):
                    content = teletype.extractText(p)
                    html += f"<p>{content}</p>\n"
    return html

def extract_text_from_xlsx(path):
    """Extraire du texte depuis un fichier XLSX"""
    path = config.FILES_DIR / path  # Utiliser config.FILES_DIR
    try:
        from openpyxl import load_workbook
        wb = load_workbook(filename=str(path), data_only=True)
        content = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for cell in row:
                    if isinstance(cell, str):
                        content.append(cell.strip())
                    elif cell is not None:
                        content.append(str(cell).strip())
        return "\n".join(content)
    except Exception as e:
        print(f"[ERREUR] Lecture XLSX Ã©chouÃ©e : {e}")
        return ""

def extract_text_from_pptx(path):
    """Extraire du texte depuis un fichier PPTX"""
    path = config.FILES_DIR / path  # Utiliser config.FILES_DIR
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
        content = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    content.append(shape.text)
        return "\n".join(content)
    except Exception as e:
        print(f"[ERREUR] Lecture PPTX Ã©chouÃ©e : {e}")
        return ""


def extract_text_from_odt(path):
    """
    Extraction de texte depuis un ODT via lecture directe du ZIP + iterparse.
    Approche non récursive : évite les RecursionError sur les gros fichiers.
    """
    import zipfile
    import xml.etree.ElementTree as ET
    import config

    try:
        full_path = config.FILES_DIR / path
        text_ns = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"
        texts = []

        with zipfile.ZipFile(str(full_path), 'r') as z:
            if 'content.xml' not in z.namelist():
                return ""
            with z.open('content.xml') as f:
                # iterparse est itératif — pas de récursion, peu de mémoire
                for event, elem in ET.iterparse(f, events=("end",)):
                    local = elem.tag.split("}")[-1] if "}" in elem.tag else elem.tag
                    if local in ("p", "h", "span") and elem.text and elem.text.strip():
                        texts.append(elem.text.strip())
                    # Libérer la mémoire au fur et à mesure
                    elem.clear()

        return "\n".join(texts)

    except Exception as e:
        print(f"[ERREUR] Lecture ODT : {e}")
        return ""

def odf_to_html(path):
    """Convertir un fichier ODT en HTML"""
    from odf.opendocument import load
    from odf import text, teletype
    import config

    path = config.FILES_DIR / path  # Utiliser config.FILES_DIR pour obtenir le chemin correct
    doc = load(str(path))

    # Structure HTML de base
    html = "<html><head><meta charset='UTF-8'><title>Document ODT</title></head><body>"

    # Traitement des titres (H)
    for elem in doc.getElementsByType(text.H):
        level = elem.getAttribute("outlinelevel")
        content = teletype.extractText(elem)
        html += f"<h{level}>{content}</h{level}>\n"
    
    # Traitement des paragraphes (P)
    for elem in doc.getElementsByType(text.P):
        content = teletype.extractText(elem)
        html += f"<p>{content}</p>\n"

    html += "</body></html>"  # Fermeture des balises HTML

    return html



